"""Importación/exportación con archivos FICTICIOS generados en la propia prueba."""
import io

import pandas as pd
from openpyxl import load_workbook

TRADE_HEADER = ["País_reportante", "Flujo_comercial", "Año", "Código_SA6", "Categoría_calzado",
                "Socio_comercial", "Valor_miles_USD", "Cantidad_pares", "Valor_USD", "Valor_USD_por_par"]
TRADE_ROWS = [
    # Datos FICTICIOS para pruebas.
    ["PaisFicticio", "Importación", 2022, "640299", "Calzado deportivo", "SocioA", 100, 10000, 100000, 10],
    ["PaisFicticio", "Importación", 2023, "640299", "Calzado deportivo", "SocioA", 150, 10000, 150000, 15],
    ["PaisFicticio", "Importación", 2023, "640399", "Calzado de cuero", "SocioB", 50, 2500, 50000, 20],
    ["PaisFicticio", "Exportación", 2023, "640299", "Calzado deportivo", "SocioC", 20, 1000, 20000, 20],
    ["PaisFicticio", "Importación", 2023, "640299", "Calzado deportivo", "Mundo", 200, 12500, 200000, 16],
]


def csv_file(rows, header=TRADE_HEADER, sep=","):
    frame = pd.DataFrame(rows, columns=header)
    return io.BytesIO(frame.to_csv(index=False, sep=sep).encode("utf-8"))


def xlsx_file(rows, header=TRADE_HEADER):
    buffer = io.BytesIO()
    pd.DataFrame(rows, columns=header).to_excel(buffer, index=False)
    buffer.seek(0)
    return buffer


def upload(api, data, name, kind="comercio"):
    return api.post("/api/imports", data={"kind": kind, "file": (data, name)}, content_type="multipart/form-data")


def test_trade_import_preview_then_confirm(manager):
    response = upload(manager, csv_file(TRADE_ROWS), "comercio_FICTICIO.csv")
    assert response.status_code == 201, response.json
    batch = response.json
    assert batch["status"] == "validado"
    assert batch["valid_rows"] == 5
    assert batch["summary"]["missing_columns"] == []
    assert len(batch["summary"]["preview"]) == 5
    # Nada se escribe antes de confirmar.
    assert manager.get("/api/trade/records").json["total"] == 0
    confirmed = manager.post(f"/api/imports/{batch['id']}/confirm", json={})
    assert confirmed.status_code == 200
    assert confirmed.json["added"] == 5
    # Por defecto los agregados ("Mundo") quedan fuera para evitar doble conteo.
    assert manager.get("/api/trade/records").json["total"] == 4
    assert manager.get("/api/trade/records?include_aggregates=1").json["total"] == 5


def test_reimport_same_file_blocked_and_updates_detected(manager):
    first = upload(manager, csv_file(TRADE_ROWS), "a.csv").json
    manager.post(f"/api/imports/{first['id']}/confirm", json={})
    again = upload(manager, csv_file(TRADE_ROWS), "a_copia.csv").json
    assert again["summary"]["blocking"] is True
    assert manager.post(f"/api/imports/{again['id']}/confirm", json={}).status_code == 422

    changed = [r[:] for r in TRADE_ROWS]
    changed[0][8] = 110000  # nuevo valor USD para la misma clave
    changed[0][6] = 110
    changed[0][9] = 11
    batch = upload(manager, csv_file(changed), "b.csv").json
    result = manager.post(f"/api/imports/{batch['id']}/confirm", json={}).json
    assert result["updated"] == 1
    assert result["skipped"] == 4
    assert result["added"] == 0


def test_trade_validation_errors(manager):
    rows = [
        ["PaisFicticio", "Trueque", 2023, "640299", "X", "SocioA", 1, 1, 1000, 1000],   # flujo no válido
        ["PaisFicticio", "Importación", 1800, "640299", "X", "SocioA", 1, 1, 1000, 1000],  # año
        ["PaisFicticio", "Importación", 2023, "64AB", "X", "SocioA", 1, 1, 1000, 1000],   # SA6
        ["PaisFicticio", "Importación", 2023, "640200", "X", "SocioA", -5, 1, 1000, 1000],  # negativo
        ["PaisFicticio", "Importación", 2023, "640300", "X", "SocioZ", 1, 100, 5000, 50],  # unidades inconsistentes
    ]
    batch = upload(manager, csv_file(rows), "errores.csv").json
    assert batch["rejected"] == 4
    assert batch["valid_rows"] == 1
    messages = " ".join(i["message"] for i in batch["issues"])
    assert "Flujo 'Trueque'" in messages
    assert "Año" in messages
    assert "SA6" in messages
    assert "negativo" in messages
    warnings = [i for i in batch["issues"] if i["severity"] == "advertencia"]
    assert any("Valor_miles_USD" in w["message"] for w in warnings)


def test_duplicates_policy(manager):
    exact = TRADE_ROWS[0]
    conflicting = exact[:]
    conflicting[8] = 999999
    batch = upload(manager, csv_file([exact, exact, TRADE_ROWS[2]]), "dup.csv").json
    assert batch["valid_rows"] == 2  # el duplicado exacto se omite con advertencia
    batch = upload(manager, csv_file([exact, conflicting]), "conflicto.csv").json
    assert batch["valid_rows"] == 0 and batch["rejected"] == 2  # misma clave, valores distintos → revisión


def test_missing_columns(manager):
    batch = upload(manager, csv_file([["x", "y"]], header=["Columna1", "Columna2"]), "malo.csv").json
    assert "pais_reportante" in batch["summary"]["missing_columns"]
    assert manager.post(f"/api/imports/{batch['id']}/confirm", json={}).status_code == 422


def test_xlsx_and_semicolon_csv(manager):
    assert upload(manager, xlsx_file(TRADE_ROWS), "comercio.xlsx").json["valid_rows"] == 5
    assert upload(manager, csv_file(TRADE_ROWS, sep=";"), "comercio_pyc.csv").json["valid_rows"] == 5


def test_rejects_disguised_files(manager):
    assert upload(manager, io.BytesIO(b"PK\x03\x04no es excel"), "falso.xlsx").status_code == 415
    assert upload(manager, io.BytesIO(b"MZ\x00\x00binario"), "falso.csv").status_code == 415
    assert upload(manager, io.BytesIO(b"a,b"), "macro.xlsm").status_code == 415


def test_value_per_pair_and_summary(manager):
    batch = upload(manager, csv_file(TRADE_ROWS), "c.csv").json
    manager.post(f"/api/imports/{batch['id']}/confirm", json={})
    summary = manager.get("/api/trade/summary").json
    imports_2023 = next(r for r in summary["by_year_flow"] if r["year"] == 2023 and r["flow"] == "Importación")
    assert imports_2023["value_usd"] == 200000  # sin "Mundo"
    per_pair = next(r for r in summary["value_per_pair_by_year"] if r["year"] == 2023 and r["flow"] == "Importación")
    assert per_pair["value_per_pair"] == 16.0  # (150000+50000)/(10000+2500)
    assert summary["top_partners"][0]["partner"] == "SocioA"
    variation = next(v for v in summary["annual_variation"] if v["flow"] == "Importación")
    assert variation["value_change_pct"] == 100.0  # 100k → 200k
    filtered = manager.get("/api/trade/summary?year_from=2023&flow=Exportación").json
    assert filtered["records"] == 1
    synthesis = manager.get("/api/trade/synthesis").json
    assert synthesis["available"] is True
    assert len(synthesis["findings"]) == 3
    assert len(synthesis["limitations"]) == 2


def test_catalog_import_creates_products_and_stock(manager):
    header = ["codigo_producto", "marca", "modelo", "categoria", "sku", "talla", "color", "costo", "precio",
              "stock_minimo", "cantidad_entrada"]
    rows = [["IMP-1", "MarcaFicticia", "Modelo Importado", "Urbano", "IMP-1-40-N", "40", "Negro", "50", "90", "1", "3"],
            ["IMP-1", "MarcaFicticia", "Modelo Importado", "Urbano", "IMP-1-41-N", "41", "Negro", "50", "90", "1", ""]]
    batch = upload(manager, csv_file(rows, header=header), "catalogo.csv", kind="catalogo").json
    assert batch["valid_rows"] == 2, batch["issues"]
    result = manager.post(f"/api/imports/{batch['id']}/confirm", json={}).json
    assert result["added"] == 2
    stock = manager.get("/api/inventory/stock?q=IMP-1").json
    assert stock["summary"]["pairs"] == 3
    movements = manager.get("/api/inventory/movements").json["items"]
    assert movements[0]["reference"] == f"IMPORT-{batch['id']}"
    assert manager.get("/api/inventory/reconciliation").json["ok"] is True


def test_catalog_duplicate_sku_rejected(manager):
    header = ["codigo_producto", "marca", "modelo", "sku", "talla", "color"]
    rows = [["X1", "M", "Mod", "DUP", "40", "Rojo"], ["X1", "M", "Mod", "DUP", "41", "Rojo"]]
    batch = upload(manager, csv_file(rows, header=header), "dup.csv", kind="catalogo").json
    assert batch["valid_rows"] == 0


def test_imports_require_permission(operator):
    assert upload(operator, csv_file(TRADE_ROWS), "x.csv").status_code == 403


# ---------- Exportación ----------

def test_export_stock_csv_and_xlsx(manager, variant):
    manager.post("/api/inventory/movements", json={"variant_id": variant["id"], "movement_type": "entrada_compra",
                                                   "quantity": 4})
    response = manager.get("/api/exports/existencias?format=csv")
    assert response.status_code == 200
    text = response.data.decode("utf-8-sig")
    assert "TEST-001-40-NEG" in text
    response = manager.get("/api/exports/existencias?format=xlsx&size=40")
    workbook = load_workbook(io.BytesIO(response.data))
    assert workbook.sheetnames == ["Datos", "Metadatos"]
    rows = list(workbook["Datos"].values)
    assert rows[1][4] == "TEST-001-40-NEG"
    meta = {r[0]: r[1] for r in workbook["Metadatos"].values}
    assert "size=40" in meta["Filtros"]
    # Filtro sin resultados: no se genera un archivo vacío.
    assert manager.get("/api/exports/existencias?format=csv&size=99").status_code == 404


def test_export_filters_trade(manager):
    batch = upload(manager, csv_file(TRADE_ROWS), "t.csv").json
    manager.post(f"/api/imports/{batch['id']}/confirm", json={})
    frame = pd.read_csv(io.BytesIO(manager.get("/api/exports/comercio?format=csv&flow=Exportación").data))
    assert list(frame["Socio_comercial"]) == ["SocioC"]
    frame = pd.read_csv(io.BytesIO(manager.get("/api/exports/comercio?format=csv&year_from=2023&year_to=2023").data))
    assert set(frame["Año"]) == {2023}
    assert "Mundo" not in set(frame["Socio_comercial"])


def test_export_formula_injection_neutralized(manager, brand):
    manager.post("/api/products", json={"code": "=HYPERLINK(\"x\")", "brand_id": brand["id"], "model_name": "M"})
    text = manager.get("/api/exports/productos?format=csv").data.decode("utf-8-sig")
    assert "'=HYPERLINK" in text


def test_export_logged_and_permissions(manager, viewer, variant):
    manager.post("/api/inventory/movements", json={"variant_id": variant["id"], "movement_type": "entrada_compra",
                                                   "quantity": 1})
    manager.get("/api/exports/movimientos?format=csv")
    history = manager.get("/api/exports").json["items"]
    assert history[0]["dataset"] == "movimientos" and history[0]["row_count"] == 1
    assert viewer.get("/api/exports/movimientos?format=csv").status_code == 403


def test_pdf_reports(manager, variant):
    assert manager.get("/api/reports/comercio.pdf").status_code == 404  # sin datos comerciales
    response = manager.get("/api/reports/inventario.pdf")
    assert response.status_code == 200
    assert response.data.startswith(b"%PDF")
    batch = upload(manager, csv_file(TRADE_ROWS), "p.csv").json
    manager.post(f"/api/imports/{batch['id']}/confirm", json={})
    assert manager.get("/api/reports/comercio.pdf").data.startswith(b"%PDF")


def test_duplicate_policy_separate_operations(manager):
    first = TRADE_ROWS[0]
    second = first[:]
    second[6], second[8], second[9] = 50, 50000, 5  # misma clave, otra operación (FICTICIA)
    rejected = upload(manager, csv_file([first, second]), "dup_rechazar.csv").json
    assert rejected["valid_rows"] == 0
    data = {"kind": "comercio", "duplicate_policy": "separados", "file": (csv_file([first, second]), "dup_sep.csv")}
    batch = manager.post("/api/imports", data=data, content_type="multipart/form-data").json
    assert batch["valid_rows"] == 2 and batch["summary"]["duplicate_policy"] == "separados"
    result = manager.post(f"/api/imports/{batch['id']}/confirm", json={}).json
    assert result["added"] == 2
    records = manager.get("/api/trade/records").json["items"]
    assert sorted(r["occurrence"] for r in records) == [1, 2]
    summary = manager.get("/api/trade/summary").json
    assert summary["by_year_flow"][0]["value_usd"] == 150000  # se suman como operaciones distintas
    bad = manager.post("/api/imports", data={"kind": "comercio", "duplicate_policy": "inventada",
                                             "file": (csv_file([first]), "x.csv")}, content_type="multipart/form-data")
    assert bad.status_code == 422


def test_truncated_headers_are_explained(manager):
    header = [h[:-1] if any(c in h for c in "íñóá") else h for h in TRADE_HEADER]  # simula el defecto real
    batch = upload(manager, csv_file([TRADE_ROWS[0]], header=header), "truncado.csv").json
    assert batch["valid_rows"] == 0
    assert any("perdido su último carácter" in i["message"] for i in batch["issues"])


def test_wrong_kind_suggests_correct_one(manager):
    batch = upload(manager, csv_file(TRADE_ROWS), "comercio.csv", kind="catalogo").json
    assert batch["valid_rows"] == 0
    assert batch["summary"]["suggested_kind"] == "comercio"
    assert "Comercio exterior" in batch["issues"][0]["message"]


def test_auto_kind_detection(manager):
    data = {"kind": "auto", "file": (csv_file(TRADE_ROWS), "auto.csv")}
    batch = manager.post("/api/imports", data=data, content_type="multipart/form-data").json
    assert batch["kind"] == "comercio" and batch["summary"]["kind_detected"] is True
    header = [h[:-1] if any(c in h for c in "íñóá") else h for h in TRADE_HEADER]
    truncated = manager.post("/api/imports", data={"kind": "catalogo", "file": (csv_file([TRADE_ROWS[0]], header=header), "t.csv")},
                             content_type="multipart/form-data").json
    assert truncated["summary"]["suggested_kind"] == "comercio"
    unknown = manager.post("/api/imports", data={"kind": "auto", "file": (csv_file([["a", "b"]], header=["x", "y"]), "u.csv")},
                           content_type="multipart/form-data")
    assert unknown.status_code == 422
